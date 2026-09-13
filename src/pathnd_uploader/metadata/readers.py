"""Loads metadata manifests — always a manifest, never a per-slide sidecar
file. Real institutional exports (BDR's CSV, Mount Sinai/PART's xlsx) are
each one spreadsheet covering many slides; there's no supported format for
a standalone `slide.json` next to `slide.svs`. For a single slide, use a
manifest with one row (see `read_single_record`).
"""

from __future__ import annotations

import csv
import json
from pathlib import Path


def read_single_record(path: Path, *, sheet: str | int | None = None) -> dict:
    """Reads a manifest expected to hold exactly one slide's record — for
    validating/uploading a single slide via a small manifest (a one-row CSV
    or xlsx, or a JSON array with one object) rather than a per-slide
    sidecar file.
    """
    records = read_manifest(path, sheet=sheet)
    if len(records) != 1:
        raise ValueError(f"{path} must contain exactly one record for a single-slide command, found {len(records)}")
    return records[0]


def read_manifest(path: Path, *, sheet: str | int | None = None) -> list[dict]:
    """Reads a batch manifest: CSV, JSON-Lines, or Excel, one row/line per slide.

    Each record must include enough to locate the slide file itself — by
    convention the schema's `slide_paths` field, interpreted relative to the
    manifest's own directory unless it is absolute.

    `sheet` selects a worksheet for `.xlsx` files by name or 0-based index
    (default: the first sheet) — real institutional exports commonly ship
    multiple sheets (e.g. slide-level vs. case-level data) in one workbook.
    """
    path = Path(path)
    if path.suffix.lower() in (".jsonl", ".ndjson"):
        with path.open(encoding="utf-8") as f:
            return [json.loads(line) for line in f if line.strip()]
    if path.suffix.lower() == ".json":
        with path.open(encoding="utf-8") as f:
            data = json.load(f)
        if not isinstance(data, list):
            raise ValueError(f"{path} must contain a JSON array of per-slide records")
        return data
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8") as f:
            return list(csv.DictReader(f))
    if path.suffix.lower() == ".xlsx":
        return _read_xlsx(path, sheet=sheet)
    raise ValueError(f"Unsupported manifest format: {path.suffix} (use .csv, .json, .jsonl, or .xlsx)")


def _read_xlsx(path: Path, *, sheet: str | int | None) -> list[dict]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet is None:
        worksheet = workbook[workbook.sheetnames[0]]
    elif isinstance(sheet, int):
        worksheet = workbook[workbook.sheetnames[sheet]]
    else:
        worksheet = workbook[sheet]

    rows = worksheet.iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(rows)]
    return [
        {col: value for col, value in zip(header, row) if col}
        for row in rows
        if any(v is not None for v in row)
    ]


def resolve_slide_path(record: dict, *, manifest_dir: Path) -> Path:
    """Resolves a manifest record's declared `slide_paths` to an actual path
    on disk, relative to the manifest's own directory if not absolute.
    """
    declared = record.get("slide_paths")
    if not declared:
        raise KeyError("record has no 'slide_paths' field")
    declared_path = Path(declared)
    return declared_path if declared_path.is_absolute() else manifest_dir / declared_path
